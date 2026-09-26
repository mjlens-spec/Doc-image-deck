# -*- coding: utf-8 -*-
"""End-to-end smoke test of everything except PowerPoint and Codex, on a synthetic slide.

Usage: <runtime python> tests/pipeline_smoke.py
Runs on macOS and Windows (CI: windows-latest after `deck setup --no-calib`). Steps:
  whiteboard deck from an outline -> a synthetic image page (Noto Serif SC + Noto Sans CJK SC) saved inside a PDF ->
  page extraction from the PDF -> OCR -> corpus -> layers.py (fitting + LaMa) -> build_pptx.py -> every string is
  editable text in the PPTX. On Windows also checks that the fonts are installed for the current user.
"""
import os, sys, json, shutil, subprocess, tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
S = os.path.join(ROOT, 'scripts')
sys.path.insert(0, S)
import hostos

TITLE, BULLETS = '内容翻倍，平均互动却下降一半', ['七成参与达人粉丝不足 1,000', '高互动笔记集中在腰部达人']
ENV = hostos.child_env()


def step(name, cmd):
    print('==', name, flush=True)
    r = subprocess.run([sys.executable] + cmd, env=ENV, capture_output=True, text=True, encoding='utf-8', errors='replace')
    print((r.stdout or '')[-1500:])
    if r.returncode:
        print((r.stderr or '')[-3000:])
        raise SystemExit('%s failed (%d)' % (name, r.returncode))
    return r.stdout


def main():
    from PIL import Image, ImageDraw, ImageFont
    d = tempfile.mkdtemp(prefix='did_smoke_')
    try:
        proj = os.path.join(d, 'proj'); wd = os.path.join(proj, '00_白板稿'); os.makedirs(wd)
        outline = {'title': '测试', 'storyline': {'thesis': '冒烟测试的结论'},
                   'pages': [{'id': 'p01', 'kind': 'content', 'chapter': '测试', 'title': TITLE,
                              'blocks': [{'type': 'bullets', 'items': BULLETS}], 'notes': '讲稿。'}]}
        json.dump(outline, open(os.path.join(wd, 'outline.json'), 'w', encoding='utf-8'), ensure_ascii=False)
        json.dump({'name': '冒烟测试', 'suffix': 'AC_0925A'}, open(os.path.join(proj, 'project.json'), 'w', encoding='utf-8'), ensure_ascii=False)
        step('humanize accept', [os.path.join(S, 'humanize.py'), proj, 'accept', '--note', 'smoke'])
        step('whiteboard', [os.path.join(S, 'whiteboard.py'), proj])
        assert os.path.exists(os.path.join(wd, '冒烟测试_白板稿_AC_0925A.pptx'))

        # a synthetic "generated" slide: serif title, sans bullets, flat background, one photo-like block
        W, H = 3344, 1882
        im = Image.new('RGB', (W, H), (246, 241, 232))
        dr = ImageDraw.Draw(im)
        dr.rectangle((2300, 300, 3100, 1500), fill=(60, 90, 80))
        serif = ImageFont.truetype(hostos.find_font('NotoSerifSC-wght.ttf'), 150)
        serif.set_variation_by_axes([600])
        sans = ImageFont.truetype(hostos.find_font('NotoSansCJKsc-Regular.otf'), 70)
        dr.text((200, 300), TITLE, font=serif, fill=(11, 74, 60))
        for k, b in enumerate(BULLETS):
            dr.text((220, 800 + k * 160), b, font=sans, fill=(58, 54, 50))
        pdf = os.path.join(d, 'deck.pdf')
        im.save(pdf, resolution=250.8)                   # 3344 px on a 13.333 in page
        work = os.path.join(d, 'work')
        step('extract pages from PDF', [os.path.join(S, 'editable', 'extract_pages.py'), pdf, work])
        page = os.path.join(work, '01_原图', 'p01.png')
        assert Image.open(page).size[0] >= 3000, Image.open(page).size

        lines = json.loads(hostos.run_ocr([page]) or '[]')
        text = ''.join(L['text'] for L in lines).replace(' ', '')
        print('OCR (%s):' % hostos.ocr_backend(), [L['text'] for L in lines])
        assert TITLE.replace(' ', '') in text, 'title not read by OCR'
        assert all(len(L['chars']) == len(L['text']) for L in lines), 'character boxes do not line up'

        step('corpus', [os.path.join(S, 'editable', 'corpus.py'), work, '--outline', os.path.join(wd, 'outline.json')])
        json.dump({'pages': {}, 'replace': {}}, open(os.path.join(work, 'config.json'), 'w', encoding='utf-8'))
        shutil.copy(os.path.join(ROOT, 'tests', 'fixtures', 'calibration_test.json'), os.path.join(work, 'calibration.json'))
        step('layers (OCR vote, font fitting, LaMa)', [os.path.join(S, 'editable', 'layers.py'), work, 'p01'])
        assert os.path.exists(os.path.join(work, '02_分层', 'p01', 'layers.json'))
        assert os.path.exists(os.path.join(work, '02_分层', 'p01', 'plate.jpg'))
        out = os.path.join(d, 'editable.pptx')
        step('build pptx', [os.path.join(S, 'editable', 'build_pptx.py'), work, out])
        from pptx import Presentation
        texts = [t for sh in Presentation(out).slides[0].shapes if sh.has_text_frame
                 for p in sh.text_frame.paragraphs for t in p.text.split('\x0b')]      # \x0b: line break inside a paragraph
        print('editable text:', texts)
        for s in [TITLE] + BULLETS:
            assert any(s.replace(' ', '') == t.replace(' ', '') for t in texts), 'missing editable text: ' + s
        if hostos.IS_WIN:
            r = subprocess.run([sys.executable, os.path.join(S, 'fontsetup.py'), os.path.join(hostos.RUNTIME, 'fonts'), '--check'],
                               env=ENV)
            assert r.returncode == 0, 'fonts are not installed for the current user'
        print('SMOKE OK: %d strings editable' % (1 + len(BULLETS)))
    finally:
        shutil.rmtree(d, ignore_errors=True)


if __name__ == '__main__':
    main()

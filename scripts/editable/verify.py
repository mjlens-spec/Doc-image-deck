# -*- coding: utf-8 -*-
"""Step 5 · Render the editable deck with PowerPoint and compare every slide with its source page.

Usage: verify.py <workdir> <deck.pptx> <out.pdf> [--ref source.pdf] [pid ...]
--ref: compare against the pages of this PDF (the source deck as rendered, including pasted pictures)
instead of the bare page images.
Writes <workdir>/03_QA/cmp_<pid>.jpg (source | PowerPoint render | amplified difference), diff.json,
and checks that every converted line exists as real text in the PPTX.
"""
import os, sys, json, subprocess, glob, re
import numpy as np
from PIL import Image
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from pptrender import ppt_to_pdf
WORK, deck, pdf = os.path.abspath(sys.argv[1]), sys.argv[2], sys.argv[3]
rest = sys.argv[4:]
REF = None
if '--ref' in rest:
    k = rest.index('--ref'); REF = rest[k + 1]; rest = rest[:k] + rest[k + 2:]
man = json.load(open(os.path.join(WORK, 'manifest.json')))
pids = rest or [p['pid'] for p in man['pages']]
qa = os.path.join(WORK, '03_QA'); os.makedirs(qa, exist_ok=True)
if not os.path.exists(pdf) or os.path.getmtime(pdf) < os.path.getmtime(deck):
    ppt_to_pdf(deck, pdf)
rd = os.path.join(WORK, 'tmp', 'render'); os.makedirs(rd, exist_ok=True)
for f in glob.glob(os.path.join(rd, 'r-*.png')): os.remove(f)
subprocess.run(['pdftoppm', '-png', '-scale-to-x', '1920', '-scale-to-y', '1080', pdf, os.path.join(rd, 'r')], check=True)
rs = sorted(glob.glob(os.path.join(rd, 'r-*.png')), key=lambda f: int(re.findall(r'-(\d+)\.png', f)[0]))
refs = {}
if REF:
    for f in glob.glob(os.path.join(rd, 'o-*.png')): os.remove(f)
    subprocess.run(['pdftoppm', '-png', '-scale-to-x', '1920', '-scale-to-y', '1080', REF, os.path.join(rd, 'o')], check=True)
    ro = sorted(glob.glob(os.path.join(rd, 'o-*.png')), key=lambda f: int(re.findall(r'-(\d+)\.png', f)[0]))
    allp = [p['pid'] for p in man['pages']]
    refs = {pid: ro[allp.index(pid)] for pid in pids if allp.index(pid) < len(ro)}
res = {}
for pid, rp in zip(pids, rs):
    a = np.array(Image.open(rp).convert('RGB')).astype(np.int16)
    src = refs.get(pid) or os.path.join(WORK, '01_原图', pid + '.png')
    b = np.array(Image.open(src).convert('RGB').resize((1920, 1080), Image.LANCZOS)).astype(np.int16)
    d = np.abs(a - b).max(-1)
    res[pid] = dict(mean=round(float(d.mean()), 2), pct_over_40=round(float((d > 40).mean() * 100), 2))
    side = np.concatenate([b, a, np.stack([np.clip(d * 3, 0, 255)] * 3, -1)], 1).astype(np.uint8)
    Image.fromarray(side).resize((2880, 540), Image.LANCZOS).save(os.path.join(qa, 'cmp_%s.jpg' % pid), quality=85)
    print('%s mean diff %.2f | pixels off >40: %.2f%%' % (pid, res[pid]['mean'], res[pid]['pct_over_40']), flush=True)
# every converted line must exist as editable text
from pptx import Presentation
prs = Presentation(deck)
missing = {}
for pid, s in zip(pids, prs.slides):
    blob = ''.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame).replace('\n', '').replace('\x0b', '')
    lj = os.path.join(WORK, '02_分层', pid, 'layers.json')
    if not os.path.exists(lj): continue
    for p in json.load(open(lj))['paragraphs']:
        for l in p['lines']:
            if l['text'] not in blob:
                missing.setdefault(pid, []).append(l['text'])
json.dump(dict(diff=res, missing_text=missing), open(os.path.join(qa, 'diff.json'), 'w'), ensure_ascii=False, indent=1)
print('lines missing as editable text:', missing or 'none')

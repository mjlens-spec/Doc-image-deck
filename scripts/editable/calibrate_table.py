# -*- coding: utf-8 -*-
"""Calibration, part 2 · Baseline position table for exact line spacing in PowerPoint.

PowerPoint places the first baseline of a top-anchored, zero-margin text box at a distance from the box top
that depends only on the line pitch (not on the font size or on these two font families) and moves in
~1 pt steps. This sweeps pitches 6–120 pt in 0.25 pt steps with two sizes and both families, exports with
PowerPoint and measures the baseline of 'HH'. Result: calibration.json["baseline_table"] {pitch: offset_pt}.
Usage: calibrate_table.py <dir containing calibration.json>
"""
import os, sys, json, subprocess, glob, re
import numpy as np
from PIL import Image
Image.MAX_IMAGE_PIXELS = None
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from lxml import etree
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(1, os.path.dirname(os.path.dirname(os.path.abspath(__file__)))); import hostos as H
from pptrender import ppt_to_pdf

work = os.path.abspath(sys.argv[1])
tmp = os.path.join(work, 'tmp', 'calib_table'); os.makedirs(tmp, exist_ok=True)
prs = Presentation(); prs.slide_width, prs.slide_height = 12192000, 6858000
probes = []
s = None; x = y = 10.0; rowh = 0.0
for p in np.arange(6, 120.01, 0.25):
    for k, face in ((1.25, 'Noto Serif SC'), (1.7, 'Noto Sans CJK SC')):
        size = round(p / k * 2) / 2
        if size < 4:
            continue
        w = size * 1.6 + 6
        if s is None or x + w > 950:
            x = 10; y += rowh + 6; rowh = 0
        if s is None or y + p * 1.3 > 530:
            s = prs.slides.add_slide(prs.slide_layouts[6]); x = 10; y = 10; rowh = 0
        top = float(int(y))
        b = s.shapes.add_textbox(Emu(int(round(x * 12700))), Emu(int(round(top * 12700))), Emu(int(w * 12700)), Emu(int(p * 12700)))
        tf = b.text_frame; tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.NONE
        for a in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
            setattr(tf, a, 0)
        bp = tf._txBody.find(qn('a:bodyPr')); bp.set('anchor', 't'); bp.set('wrap', 'none')
        para = tf.paragraphs[0]
        ln = etree.SubElement(para._p.get_or_add_pPr(), qn('a:lnSpc')); etree.SubElement(ln, qn('a:spcPts')).set('val', str(int(round(p * 100))))
        r = para.add_run(); r.text = 'HH'; r.font.size = Pt(size)
        rPr = r._r.get_or_add_rPr()
        for t in ('a:latin', 'a:ea', 'a:cs'):
            etree.SubElement(rPr, qn(t)).set('typeface', face)
        probes.append(dict(slide=len(prs.slides), pitch=float(p), top=top, x=x, w=w))
        x += w + 4; rowh = max(rowh, p * 1.3)
pp = os.path.join(tmp, 'table.pptx'); prs.save(pp)
pdf = ppt_to_pdf(pp, os.path.join(tmp, 'table.pdf'))
DPI = 600; S = DPI / 72.0
H.render_pdf(pdf, os.path.join(tmp, 'c'), dpi=DPI)
pages = sorted(glob.glob(os.path.join(tmp, 'c-*.png')), key=lambda f: int(re.findall(r'-(\d+)\.png', f)[0]))
cur, im = None, None
offs = {}
for pr in probes:
    if pr['slide'] != cur:
        cur = pr['slide']; im = np.array(Image.open(pages[cur - 1]).convert('L')) < 128
    y0 = int((pr['top'] - 3) * S); y1 = int((pr['top'] + pr['pitch'] * 1.3) * S)
    band = im[y0:y1, int((pr['x'] + 0.5) * S):int((pr['x'] + pr['w'] - 1) * S)]
    rows = np.where(band.any(1))[0]
    if len(rows):
        offs.setdefault(pr['pitch'], []).append((y0 + rows[-1] + 1) / S - pr['top'])
table = {}
for p, v in sorted(offs.items()):
    table['%.2f' % p] = float(np.median(v)) if len(v) != 2 or abs(v[0] - v[1]) <= 0.3 else float(min(v, key=lambda o: abs(o - 0.752 * p)))
cp = os.path.join(work, 'calibration.json')
cal = json.load(open(cp)) if os.path.exists(cp) else {}
cal['baseline_table'] = table
cal['baseline_rule'] = 'exact line spacing: baseline below text-box top = table[pitch rounded to 0.25 pt]'
json.dump(cal, open(cp, 'w'), indent=1)
print('baseline table: %d pitches' % len(table))

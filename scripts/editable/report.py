# -*- coding: utf-8 -*-
"""Step 10 helper · Numbers for the delivery check: per-page conversion counts, what stayed in the plate
and why, fonts used, OCR corrections, and places where the page wording differs from the draft copy.
Usage: report.py <workdir> <deck.pptx>   -> prints Markdown tables; also writes 03_QA/stats.json"""
import os, sys, json, collections
from pptx import Presentation
from pptx.oxml.ns import qn

work, deck = sys.argv[1], sys.argv[2]
man = json.load(open(os.path.join(work, 'manifest.json')))
WHY = {'excluded': '排除区（截图 / 海报 / 表情包等）', 'symbol': '单个符号', 'ornament': '装饰符号', 'garbled': '截图乱码',
       'too-small': '字高不足阈值', 'low-contrast': '对比度过低', 'no-ink': '未取到墨迹', 'no-fit': '字体未拟合'}
rows, why_tot, style_tot = [], collections.Counter(), collections.Counter()
fixes_tot = collections.Counter(); drafts = []
for pg in man['pages']:
    lj = os.path.join(work, '02_分层', pg['pid'], 'layers.json')
    if not os.path.exists(lj):
        continue
    L = json.load(open(lj))
    nl = sum(len(p['lines']) for p in L['paragraphs'])
    w = collections.Counter()
    for s in L['skipped']:
        k = s['why'].split('(')[0]
        k = 'busy' if k.startswith('busy') else k
        w[k] += 1; why_tot[k] += 1
    for p in L['paragraphs']:
        style_tot[(p['family'], p['weight'])] += len(p['lines'])
    for f in L['corrections']:
        for v in (f.get('via') or 'other').split('+'):
            fixes_tot[v] += 1
    for d in L.get('draft_diffs', []):
        for r in d['rejected']:
            drafts.append((pg['pid'], d['text'], r['page'], r['draft']))
    rows.append((pg['pid'], len(L['paragraphs']), nl, sum(w.values()), dict(w)))
prs = Presentation(deck)
fonts = collections.Counter()
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for e in sh._element.iter(qn('a:ea')):
                fonts[e.get('typeface')] += 1
print('| 页 | 文本框 | 可编辑行 | 留在底图 |\n|---|---:|---:|---:|')
for r in rows:
    print('| %s | %d | %d | %d |' % (r[0].upper(), r[1], r[2], r[3]))
print('\n合计：%d 个文本框，%d 行可编辑文字，%d 处留在底图' % (sum(r[1] for r in rows), sum(r[2] for r in rows), sum(r[3] for r in rows)))
print('\n留在底图的原因：', {WHY.get(k, '照片上的文字' if k == 'busy' else k): v for k, v in why_tot.most_common()})
print('字体 × 字重（行数）：', {'%s %s' % k: v for k, v in style_tot.most_common()})
print('PPTX 中文字体名（run 数）：', dict(fonts))
print('OCR 纠正来源：', dict(fixes_tot))
print('\n页面与旧稿不一致（保留页面原文）：')
for d in drafts:
    print('- %s「%s」：页面为「%s」，旧稿为「%s」' % (d[0].upper(), d[1][:30], d[2] or '（无）', d[3] or '（无）'))
json.dump(dict(pages=rows, skipped=why_tot, styles={'%s %s' % k: v for k, v in style_tot.items()}, fonts=fonts,
               fixes=fixes_tot, draft_diffs=drafts), open(os.path.join(work, '03_QA', 'stats.json'), 'w'), ensure_ascii=False, indent=1)

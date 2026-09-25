# -*- coding: utf-8 -*-
"""Step 4 · Assemble the editable deck.

Starts from the source PPTX (so logos, page numbers and any pasted pictures keep their exact
position, order and formatting), swaps each full-bleed page image for its clean plate and inserts
the fitted text boxes directly above the plate. For a PDF source, a blank 16:9 deck is used.

Usage: build_pptx.py <workdir> <out.pptx> [pid ...]
"""
import os, re, sys, io, json, copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from lxml import etree
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
sys.path.insert(1, os.path.dirname(HERE)); import hostos
from fonts import face_of, render, font
from layers_spacing import track_list, char_layout

WORK = os.path.abspath(sys.argv[1]); OUT = sys.argv[2]
ONLY = set(sys.argv[3:])
CFG = json.load(open(os.path.join(WORK, 'config.json'))) if os.path.exists(os.path.join(WORK, 'config.json')) else {}
man = json.load(open(os.path.join(WORK, 'manifest.json')))
_cp = os.path.join(WORK, 'calibration.json')
if not os.path.exists(_cp):
    _cp = hostos.CALIB
cal = json.load(open(_cp))
TABLE = {float(k): v for k, v in cal['baseline_table'].items()}
SW, SH = man['slide_size']
FONT_REMAP = CFG.get('font_remap', {'Source Han Sans SC': 'Noto Sans CJK SC', 'Source Han Serif SC': 'Noto Serif SC'})

def base_offset(pitch_pt):
    q = min(TABLE, key=lambda k: abs(k - pitch_pt))
    if abs(q - pitch_pt) > 0.13:                  # outside the measured range: linear fallback
        return 0.752 * pitch_pt
    return TABLE[q]

def q25(v):
    return round(v * 4) / 4.0

def set_run(r, fam, w, size_pt, color, spc_pt, grad=None):
    face, bold = face_of(fam, w)
    r.font.size = Pt(size_pt); r.font.bold = bold
    rPr = r._r.get_or_add_rPr()
    rPr.set('lang', 'zh-CN'); rPr.set('altLang', 'en-US')
    if abs(spc_pt) >= 0.01:
        rPr.set('spc', str(int(round(spc_pt * 100))))
    if grad:
        g = etree.SubElement(rPr, qn('a:gradFill')); g.set('rotWithShape', '1')
        gl = etree.SubElement(g, qn('a:gsLst'))
        for pos, c in ((0, grad[0]), (100000, grad[1])):
            gs = etree.SubElement(gl, qn('a:gs')); gs.set('pos', str(pos))
            etree.SubElement(gs, qn('a:srgbClr')).set('val', c)
        lin = etree.SubElement(g, qn('a:lin')); lin.set('ang', '5400000'); lin.set('scaled', '0')
    else:
        r.font.color.rgb = RGBColor.from_string(color)
    for t in ('a:latin', 'a:ea', 'a:cs'):
        e = etree.SubElement(rPr, qn(t)); e.set('typeface', face)

def add_para_box(slide, para, k, PT):
    lines = para['lines']
    fam, w = para['family'], para['weight']
    size_pt = para['px'] * PT
    pitch_pt = q25(para['pitch'] * PT) if len(lines) > 1 else q25(1.25 * size_pt)
    top_pt = lines[0]['base'] * PT - base_offset(pitch_pt)
    # advance extent of each line (PowerPoint adds spc after every character)
    ext = []
    for l in lines:
        f = font(fam, w, l['px'])
        sp_, fac_ = char_layout(l)
        adv = sum(f.getlength(c) * k_ for c, k_ in zip(l['text'], fac_)) + sum(sp_)
        ext.append((l['left'], l['left'] + adv))
    al = para['align']
    if al == 'ctr':
        cx = sum((a + b) / 2 for a, b in ext) / len(ext)
        wid = max(b - a for a, b in ext) + para['px'] * 1.0
        x0 = cx - wid / 2
    elif al == 'r':
        rx = sum(b for a, b in ext) / len(ext)
        wid = max(b - a for a, b in ext) + para['px'] * 1.0
        x0 = rx - wid
    else:
        x0 = min(a for a, b in ext)
        wid = max(b for a, b in ext) - x0 + para['px'] * 1.0
    tb = slide.shapes.add_textbox(Emu(int(round(x0 * PT * 12700))), Emu(int(round(top_pt * 12700))),
                                  Emu(int(round(wid * PT * 12700))), Emu(int(round(pitch_pt * len(lines) * 12700))))
    tb.name = '文字 %02d · %s' % (k, lines[0]['text'][:12])
    tf = tb.text_frame
    tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.NONE
    for a in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, a, 0)
    bp = tf._txBody.find(qn('a:bodyPr')); bp.set('anchor', 't'); bp.set('wrap', 'none')
    p = tf.paragraphs[0]
    pPr = p._p.get_or_add_pPr()
    pPr.set('algn', {'l': 'l', 'ctr': 'ctr', 'r': 'r'}[al])
    ln = etree.SubElement(pPr, qn('a:lnSpc')); etree.SubElement(ln, qn('a:spcPts')).set('val', str(int(round(pitch_pt * 100))))
    for i, l in enumerate(lines):
        if i:
            br = etree.SubElement(p._p, qn('a:br'))
            etree.SubElement(br, qn('a:rPr')).set('sz', str(int(round(size_pt * 100))))
        tl, fac = char_layout(l)
        k = 0
        for rr in l['runs']:
            # consecutive characters with the same colour, spacing and size share one run
            parts = []
            for c in rr['text']:
                sp = tl[k] if k < len(tl) else l['track_px']; fc = fac[k] if k < len(fac) else 1.0; k += 1
                if parts and abs(parts[-1][1] - sp) < 1e-6 and parts[-1][2] == fc:
                    parts[-1][0] += c
                else:
                    parts.append([c, sp, fc])
            for part, sp, fc in parts:
                r = p.add_run(); r.text = part
                set_run(r, fam, w, l['px'] * fc * PT, rr['color'], sp * PT, l.get('gradient'))
    end = p._p.find(qn('a:endParaRPr'))
    if end is None:
        end = etree.SubElement(p._p, qn('a:endParaRPr'))
    end.set('sz', str(int(round(size_pt * 100)))); end.set('lang', 'zh-CN')
    return tb

def remap_fonts(el):
    for t in el.iter(qn('a:latin'), qn('a:ea'), qn('a:cs')):
        if t.get('typeface') in FONT_REMAP:
            t.set('typeface', FONT_REMAP[t.get('typeface')])

def build():
    if man['kind'] == 'pptx':
        prs = Presentation(man['source'])
    else:
        prs = Presentation(); prs.slide_width, prs.slide_height = SW, SH
    report = []
    pages = {p['pid']: p for p in man['pages']}
    for n, pg in enumerate(man['pages'], 1):
        pid = pg['pid']
        if man['kind'] == 'pptx':
            slide = prs.slides[n - 1]
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
        remap_fonts(slide.shapes._spTree)
        lj = os.path.join(WORK, '02_分层', pid, 'layers.json')
        if 'image' not in pg or not os.path.exists(lj) or (ONLY and pid not in ONLY):
            report.append(dict(page=pid, rebuilt=False)); continue
        L = json.load(open(lj))
        W, H = L['size']
        PT = (SW / 12700.0) / W
        plate = os.path.join(WORK, '02_分层', pid, 'plate.jpg')
        if man['kind'] == 'pptx':
            pic = [s for s in slide.shapes if s.shape_id == pg['bg_shape_id']][0]
            blip = pic._element.find('.//' + qn('a:blip'))
            old = blip.get(qn('r:embed'))
            _, rid = slide.part.get_or_add_image_part(plate)
            blip.set(qn('r:embed'), rid)
            if old != rid and slide.part._rel_ref_count(old) == 0:
                slide.part.drop_rel(old)
            pic.name = '底图（已去除文字）'
            anchor = pic._element
        else:
            pic = slide.shapes.add_picture(plate, 0, 0, SW, SH); pic.name = '底图（已去除文字）'
            anchor = pic._element
        nt = 0
        for k, para in enumerate(L['paragraphs'], 1):
            tb = add_para_box(slide, para, k, PT)
            anchor.addnext(tb._element); anchor = tb._element
            nt += 1
        report.append(dict(page=pid, rebuilt=True, text_boxes=nt, lines=sum(len(p['lines']) for p in L['paragraphs']),
                           skipped=len(L['skipped'])))
    if ONLY:                                          # test build: keep only the requested slides
        ids = prs.slides._sldIdLst
        for n, pg in reversed(list(enumerate(man['pages'], 1))):
            if pg['pid'] not in ONLY:
                sid = ids[n - 1]
                prs.part.drop_rel(sid.rId); ids.remove(sid)
    prs.core_properties.title = (prs.core_properties.title or os.path.splitext(os.path.basename(man['source']))[0]) + '（可编辑版）'
    prs.save(OUT)
    return report

if __name__ == '__main__':
    rep = build()
    json.dump(rep, open(os.path.join(WORK, 'build_report.json'), 'w'), ensure_ascii=False, indent=1)
    print('saved', OUT, '%.1f MB' % (os.path.getsize(OUT) / 1048576), '| slides', len(rep), '| rebuilt', sum(r['rebuilt'] for r in rep))

# -*- coding: utf-8 -*-
"""Step 0 (once per machine / font set) · Measure how PowerPoint places and draws each typeface.

Builds a test deck (every family x weight, several sizes, exact line spacing, letter spacing),
exports it with PowerPoint, and fits
  baseline_from_box_top = A * line_pitch + B * font_size      (per family)
and checks that every weight renders with its own stroke weight (catches broken variable-font
instances) and that advance + spc matches FreeType.
Usage: calibrate.py <workdir>      -> <workdir>/calibration.json
"""
import os, sys, json, subprocess, shutil
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from lxml import etree
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
sys.path.insert(1, os.path.dirname(os.path.dirname(os.path.abspath(__file__)))); import hostos as H
from fonts import FAMILIES, face_of, font, render
from pptrender import ppt_to_pdf

work = sys.argv[1]
cal = os.path.join(work, 'tmp', 'calib'); os.makedirs(cal, exist_ok=True)
SW, SH = 12192000, 6858000
prs = Presentation(); prs.slide_width, prs.slide_height = SW, SH
probes = []
def tb(slide, x_pt, top_pt, text, fam, w, size, pitch, spc=0.0):
    face, bold = face_of(fam, w)
    b = slide.shapes.add_textbox(Emu(int(x_pt * 12700)), Emu(int(top_pt * 12700)), Emu(int(400 * 12700)), Emu(int(pitch * 12700 * 1.2)))
    tf = b.text_frame; tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.NONE
    for a in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, a, 0)
    bp = tf._txBody.find(qn('a:bodyPr')); bp.set('anchor', 't'); bp.set('wrap', 'none')
    p = tf.paragraphs[0]
    ln = etree.SubElement(p._p.get_or_add_pPr(), qn('a:lnSpc')); etree.SubElement(ln, qn('a:spcPts')).set('val', str(int(round(pitch * 100))))
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    rPr = r._r.get_or_add_rPr()
    for t in ('a:latin', 'a:ea', 'a:cs'):
        etree.SubElement(rPr, qn(t)).set('typeface', face)
    if spc: rPr.set('spc', str(int(round(spc * 100))))
# slide per family: rows of (weight, size, pitch)
for fam in FAMILIES:
    ws = list(FAMILIES[fam]['weights'])
    rows = [(w, 20, 1.0) for w in ws] + [(ws[1], 12, 1.6), (ws[1], 30, 1.0), (ws[1], 30, 1.4), (ws[1], 44, 1.2), (ws[1], 16, 2.0)]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    top = 20.0
    for i, (w, size, k) in enumerate(rows):
        pitch = size * k
        tb(s, 20, top, 'HHHH 一二三四五六七八九十 本案设计', fam, w, size, pitch, spc=0 if i % 2 == 0 else 3.0)
        probes.append(dict(fam=fam, w=w, size=size, pitch=pitch, top=top, spc=0 if i % 2 == 0 else 3.0, slide=len(prs.slides)))
        top += pitch * 1.25 + 8
pp = os.path.join(cal, 'calib.pptx'); prs.save(pp)
pdf = ppt_to_pdf(pp, os.path.join(cal, 'calib.pdf'))
DPI = 400
for f in os.listdir(cal):
    if f.startswith('c-') and f.endswith('.png'): os.remove(os.path.join(cal, f))
H.render_pdf(pdf, os.path.join(cal, 'c'), dpi=DPI)
pages = sorted(f for f in os.listdir(cal) if f.startswith('c-') and f.endswith('.png'))
imgs = [np.array(Image.open(os.path.join(cal, f)).convert('L')) for f in pages]
S = DPI / 72.0
res = {}
for pr in probes:
    im = imgs[pr['slide'] - 1] < 128
    y0 = int((pr['top'] - 2) * S); y1 = int((pr['top'] + pr['pitch'] * 1.2 + 2) * S)
    x0 = int(20 * S)
    # the 'HHHH' block: first ~2.9 em
    hx1 = x0 + int(pr['size'] * 3.0 * S)
    band = im[y0:y1, x0:hx1]
    rows = np.where(band.any(1))[0]
    base_px = y0 + rows[-1] + 1
    pr['base_off'] = base_px / S - pr['top']                # baseline below box top (pt)
    # CJK stroke density for weight check: the '一…十 本案设计' part
    cj = im[y0:y1, hx1 + int(pr['size'] * 0.2 * S): x0 + int(pr['size'] * 21 * S)]
    pr['density'] = float(cj[:, cj.any(0)].mean()) if cj.any() else 0
    cols = np.where(im[y0:y1].any(0))[0]
    pr['width_pt'] = (cols[-1] + 1 - cols[0]) / S
    rt = render('HHHH 一二三四五六七八九十 本案设计', pr['fam'], pr['w'], pr['size'] * S, pr['spc'] * S)
    pr['width_pred_pt'] = (rt['x1'] - rt['x0']) / S
for fam in FAMILIES:
    P = [p for p in probes if p['fam'] == fam]
    X = np.array([[p['pitch'], p['size']] for p in P]); y = np.array([p['base_off'] for p in P])
    (A, B), *_ = np.linalg.lstsq(X, y, rcond=None)
    err = float(np.abs(X @ [A, B] - y).max())
    dens = {p['w']: round(p['density'], 4) for p in P if p['size'] == 20}
    wr = [round(p['width_pt'] / p['width_pred_pt'], 4) for p in P]
    res[fam] = dict(A=float(A), B=float(B), max_err_pt=err, density_by_weight=dens, width_ratio=wr)
    print(fam, 'A=%.4f B=%.4f maxerr=%.3fpt' % (A, B, err), 'density', dens, 'width ratio', min(wr), max(wr))
json.dump(dict(families=res, probes=probes), open(os.path.join(work, 'calibration.json'), 'w'), indent=1)

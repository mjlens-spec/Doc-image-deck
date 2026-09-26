# -*- coding: utf-8 -*-
"""Stage 1 · Render outline.json as a plain "whiteboard" deck: final wording, clear structure, no design.

Usage: whiteboard.py <project> [--no-punct]
Reads  <project>/00_白板稿/outline.json   (schema: references/01_白板稿.md)
Writes <project>/00_白板稿/<name>_白板稿_<suffix>.pptx and outline.md (reading copy)
Before rendering: every string must have been through the 去 AI 味 step (humanize.py, humanizer-zh), otherwise
this refuses to run; then punct.py tidies the full stops in outline.json (titles and short phrases lose them; see
标点整理.md); --no-punct skips the punctuation step.

Every slide: black text on white, thin grey rules, 微软雅黑 for Chinese and Arial for Latin/digits.
The font size of each slide is chosen so its content fits; the visual brief, layout and imagery hints go to the notes.
After rendering, visual.py checks the per-slide visual plan and writes 视觉规划.md (reported here, enforced by deck prompts).
"""
import os, sys, math, json
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import deckenv as E

EA, LATIN = 'Microsoft YaHei', 'Arial'
INK, GREY, RULE = RGBColor(0x1F, 0x1F, 0x1F), RGBColor(0x6B, 0x6B, 0x6B), RGBColor(0xBF, 0xBF, 0xBF)
W_IN, H_IN = 13.333, 7.5
ML, MR = 0.6, 0.6
CW = W_IN - ML - MR

_font = None
def text_em(s):
    """Width of s in em (CJK = 1, Latin/digits ≈ 0.55), measured with 微软雅黑 when available."""
    global _font
    if _font is None:
        import hostos
        p = hostos.find_font('msyh.ttc', 'msyh.ttf', 'Songti.ttc', 'simsun.ttc')
        _font = ImageFont.truetype(p, 100) if p else False
    if _font:
        return _font.getlength(s) / 100.0
    return sum(1.0 if ord(c) > 0x2E80 else 0.55 for c in s)

def lines_for(s, size_pt, width_in):
    per_line = max(1.0, width_in * 72.0 / size_pt)
    return max(1, sum(max(1, math.ceil(text_em(part) / per_line)) for part in s.split('\n')))

def style_run(r, size, bold=False, color=INK):
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    for tag, face in (('a:latin', LATIN), ('a:ea', EA), ('a:cs', LATIN)):
        for old in rPr.findall(qn(tag)): rPr.remove(old)
        etree.SubElement(rPr, qn(tag)).set('typeface', face)

def textbox(slide, x, y, w, h, paras, anchor='t', align=PP_ALIGN.LEFT, spacing=1.2):
    """paras: list of [(text, size, bold, color), ...] runs per paragraph"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    for a in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, a, 0)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        if i:
            p.space_before = Pt(runs[0][1] * 0.45)
        for text, size, bold, color in runs:
            r = p.add_run(); r.text = text
            style_run(r, size, bold, color)
    return tb

def rule(slide, x, y, w, color=RULE, weight=0.75):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    return ln

# ─────────────────────────────── block renderers: est(size) -> height in inches; draw(slide, y, size, h)
def blk_height(b, s, width=CW):
    t = b.get('type')
    lh = s * 1.3 / 72.0
    if t in ('bullets', 'steps'):
        items = [it if isinstance(it, str) else it.get('text', '') for it in b.get('items', [])]
        return sum(lines_for(x, s, width - 0.3) * lh + s * 0.45 / 72 for x in items)
    if t == 'numbered':
        return sum((1 + lines_for(it.get('text', '') or ' ', s, width - 0.9)) * lh + s * 0.6 / 72 for it in b.get('items', []))
    if t == 'kpis':
        return s * 2.2 * 1.25 / 72 + 2 * lh
    if t == 'table':
        n = len(b.get('header', [])) or 1
        colw = width / n
        rows = [b.get('header', [])] + b.get('rows', [])
        return sum(max(lines_for(str(c), s - 1, colw - 0.15) for c in r) * (s - 1) * 1.35 / 72 + 0.1 for r in rows) + (lh if b.get('caption') else 0)
    if t == 'columns':
        cols = b.get('columns', [])
        cw = width / max(1, len(cols)) - 0.25
        return max((lh + sum(lines_for(x, s, cw - 0.25) * lh + s * 0.4 / 72 for x in c.get('items', []))) for c in cols) if cols else 0
    if t in ('quote', 'callout', 'text', 'heading'):
        return lines_for(b.get('text', ''), s, width - 0.4) * lh + (lh if b.get('source') else 0) + 0.1
    return lh

def draw_block(slide, b, x, y, w, s):
    t = b.get('type')
    if t in ('bullets', 'steps'):
        items = [it if isinstance(it, str) else it.get('text', '') for it in b.get('items', [])]
        paras = [[(('%d. ' % (i + 1)) if t == 'steps' else '• ', s, False, GREY), (it, s, False, INK)] for i, it in enumerate(items)]
        textbox(slide, x, y, w, blk_height(b, s, w), paras)
    elif t == 'numbered':
        yy = y
        for it in b.get('items', []):
            h = (1 + lines_for(it.get('text', '') or ' ', s, w - 0.9)) * s * 1.3 / 72
            textbox(slide, x, yy, 0.8, h, [[(it.get('label', ''), s * 1.3, True, GREY)]])
            paras = [[(it.get('title', ''), s, True, INK)]]
            if it.get('text'):
                paras.append([(it['text'], s, False, INK)])
            textbox(slide, x + 0.9, yy, w - 0.9, h, paras)
            yy += h + s * 0.6 / 72
    elif t == 'kpis':
        items = b.get('items', [])
        cw = w / max(1, len(items))
        for i, it in enumerate(items):
            textbox(slide, x + i * cw, y, cw - 0.2, blk_height(b, s, w),
                    [[(it.get('value', ''), s * 2.2, True, INK)], [(it.get('label', ''), s, False, GREY)]])
    elif t == 'table':
        rows = [b.get('header', [])] + b.get('rows', [])
        n = max(len(r) for r in rows)
        weights = [max(1.0, max(text_em(str(r[j])) if j < len(r) else 1 for r in rows)) for j in range(n)]
        tot = sum(weights)
        yy = y
        if b.get('caption'):
            textbox(slide, x, yy, w, s * 1.3 / 72, [[(b['caption'], s, True, INK)]]); yy += s * 1.4 / 72
        h = blk_height(b, s, w) - (s * 1.3 / 72 if b.get('caption') else 0)
        gt = slide.shapes.add_table(len(rows), n, Inches(x), Inches(yy), Inches(w), Inches(h))
        tbl = gt.table
        tblPr = tbl._tbl.tblPr
        for child in list(tblPr):                  # plain table: no built-in theme style
            if child.tag == qn('a:tableStyleId'): tblPr.remove(child)
        etree.SubElement(tblPr, qn('a:tableStyleId')).text = '{5940675A-B579-460E-94D1-54222C63F5DA}'   # "No Style, Table Grid"
        for j in range(n):
            tbl.columns[j].width = Inches(w * weights[j] / tot)
        for i, r in enumerate(rows):
            for j in range(n):
                cell = tbl.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(0.06); cell.margin_top = cell.margin_bottom = Inches(0.03)
                tf = cell.text_frame; tf.word_wrap = True
                p = tf.paragraphs[0]; run = p.add_run(); run.text = str(r[j]) if j < len(r) else ''
                style_run(run, s - 1, i == 0, INK)
    elif t == 'columns':
        cols = b.get('columns', [])
        cw = w / max(1, len(cols))
        for i, c in enumerate(cols):
            paras = [[(c.get('title', ''), s, True, INK)]] + [[('• ', s, False, GREY), (x_, s, False, INK)] for x_ in c.get('items', [])]
            textbox(slide, x + i * cw, y, cw - 0.25, blk_height(b, s, w), paras)
    elif t in ('quote', 'callout'):
        h = blk_height(b, s, w)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.05), Inches(h - 0.1))
        bar.fill.solid(); bar.fill.fore_color.rgb = GREY; bar.line.fill.background()
        paras = [[(b.get('text', ''), s, t == 'callout', INK)]]
        if b.get('source'):
            paras.append([('—— ' + b['source'], s * 0.85, False, GREY)])
        textbox(slide, x + 0.25, y, w - 0.25, h, paras)
    elif t == 'heading':
        textbox(slide, x, y, w, blk_height(b, s, w), [[(b.get('text', ''), s * 1.05, True, INK)]])
    else:
        textbox(slide, x, y, w, blk_height(b, s, w), [[(b.get('text', ''), s, False, INK)]])

def render_page(prs, page, n, total, label=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    kind = page.get('kind', 'content')
    if kind in ('cover', 'closing', 'section'):
        y = 2.3 if kind != 'section' else 2.7
        if page.get('kicker'):
            textbox(s, ML, y - 0.55, CW, 0.4, [[(page['kicker'], 14, False, GREY)]]);
        size = 40 if kind == 'cover' else 36
        h = lines_for(page.get('title', ''), size, CW) * size * 1.25 / 72
        textbox(s, ML, y, CW, h, [[(page.get('title', ''), size, True, INK)]])
        y += h + 0.35
        rule(s, ML, y, 3.0, GREY, 1.0); y += 0.3
        if page.get('subtitle'):
            textbox(s, ML, y, CW, 0.9, [[(page['subtitle'], 20, False, INK)]]); y += 0.8
        for b in page.get('blocks', []):
            hb = blk_height(b, 16)
            draw_block(s, b, ML, y, CW, 16); y += hb + 0.15
    else:
        y = 0.35
        if page.get('kicker'):
            textbox(s, ML, y, CW, 0.3, [[(page['kicker'], 11, False, GREY)]]); y += 0.32
        th = lines_for(page.get('title', ''), 26, CW) * 26 * 1.25 / 72
        textbox(s, ML, y, CW, th, [[(page.get('title', ''), 26, True, INK)]]); y += th + 0.08
        if page.get('subtitle'):
            sh = lines_for(page['subtitle'], 15, CW) * 15 * 1.3 / 72
            textbox(s, ML, y, CW, sh, [[(page['subtitle'], 15, False, GREY)]]); y += sh + 0.05
        rule(s, ML, y + 0.08, CW); y += 0.3
        bottom = 7.0
        if page.get('footnote'):
            bottom -= 0.3
        if page.get('takeaway'):
            bottom -= 0.55
        blocks = page.get('blocks', [])
        gap = 0.22
        size = 11
        for cand in (18, 17, 16, 15, 14, 13, 12, 11):
            if sum(blk_height(b, cand) for b in blocks) + gap * max(0, len(blocks) - 1) <= bottom - y:
                size = cand; break
        for b in blocks:
            draw_block(s, b, ML, y, CW, size); y += blk_height(b, size) + gap
        if page.get('takeaway'):
            ty = (6.45 if not page.get('footnote') else 6.15)
            rule(s, ML, ty, CW)
            textbox(s, ML, ty + 0.1, CW, 0.45, [[(page['takeaway'], 15, True, INK)]])
        if page.get('footnote'):
            textbox(s, ML, 6.75, CW - 1.0, 0.3, [[(page['footnote'], 9, False, GREY)]])
    textbox(s, W_IN - MR - 1.5, 7.05, 1.5, 0.25, [[(label or '%d / %d' % (n, total), 9, False, GREY)]], align=PP_ALIGN.RIGHT)
    notes = []
    if page.get('chapter') or page.get('source'):
        src = page.get('source')
        notes.append('【章节】%s｜【来源】%s' % (page.get('chapter') or '—', '、'.join(src) if isinstance(src, list) else (src or '—')))
    v = page.get('visual') if isinstance(page.get('visual'), dict) else {}
    if v:
        notes.append('【视觉】%s｜%s｜%s｜焦点：%s｜%s' % (
            v.get('message', ''), E.STRUCTURES.get(v.get('structure'), (v.get('structure', ''),))[0], v.get('form', ''),
            v.get('focal', ''), E.SKELETONS.get(v.get('skeleton'), (v.get('skeleton', ''),))[0]))
    if page.get('layout_hint'):
        notes.append('【版面建议】' + page['layout_hint'])
    if page.get('image_hint'):
        notes.append('【配图建议】' + page['image_hint'])
    if page.get('tone'):
        notes.append('【明暗】' + page['tone'])
    if page.get('notes'):
        notes.append('【讲稿】' + page['notes'])
    if notes:
        s.notes_slide.notes_text_frame.text = '\n'.join(notes)

def outline_md(outline):
    out = ['# ' + outline.get('title', ''), '']
    for i, p in enumerate(outline['pages'], 1):
        out.append('## P%02d %s' % (i, p.get('title', '')))
        if p.get('chapter') or p.get('source'):
            src = p.get('source')
            out.append('> 章节：%s；来源：%s' % (p.get('chapter') or '—', '、'.join(src) if isinstance(src, list) else (src or '—')))
        for role, t in E.page_strings(p):
            if role != 'title':
                out.append('- %s：%s' % (role, t))
        v = p.get('visual') if isinstance(p.get('visual'), dict) else {}
        if v:
            out.append('> 视觉：%s；%s；%s；焦点：%s；构图：%s' % (
                v.get('message', ''), E.STRUCTURES.get(v.get('structure'), (v.get('structure', ''),))[0], v.get('form', ''),
                v.get('focal', ''), E.SKELETONS.get(v.get('skeleton'), (v.get('skeleton', ''),))[0]))
        for k, lab in (('layout_hint', '版面建议'), ('image_hint', '配图建议'), ('notes', '讲稿')):
            if p.get(k):
                out.append('> %s：%s' % (lab, p[k]))
        out.append('')
    return '\n'.join(out)

def render(proj):
    cfg = E.load_project(proj)
    outline = E.load_outline(proj)
    prs = Presentation(); prs.slide_width, prs.slide_height = E.SW, E.SH
    total = len(outline['pages'])
    labels = E.page_labels({'page_number': {'corner': 'br', 'skip_first': False, 'skip_last': False}}, outline['pages'])
    for n, page in enumerate(outline['pages'], 1):
        page.setdefault('id', 'p%02d' % n)
        render_page(prs, page, n, total, '附录 %s' % labels[n - 1] if page.get('kind') == 'appendix' else None)
    prs.core_properties.title = outline.get('title', cfg['name']) + '（白板稿）'
    out = os.path.join(proj, E.D_WHITE, E.out_name(cfg, '白板稿', 'pptx'))
    prs.save(out)
    open(os.path.join(proj, E.D_WHITE, 'outline.md'), 'w', encoding='utf-8').write(outline_md(outline))
    return out

def main():
    args = [x for x in sys.argv[1:] if not x.startswith('--')]
    proj = os.path.abspath(args[0])
    import humanize
    todo = humanize.pending(proj)
    if todo:
        raise SystemExit('还有 %d 条文案没有做去 AI 味处理（如 %s）。先运行 deck humanize <项目> export%s，按 humanizer-zh 处理后 import；'
                         '用户指定原话的条目用 deck humanize <项目> accept --note "说明"。'
                         % (len(todo), todo[0][0], ' --changed' if humanize.load_done(proj) else ''))
    import storyline
    serrs, swarns, spath = storyline.run(proj)
    for w in swarns:
        print('故事线提示：' + w)
    if serrs:
        raise SystemExit('故事线未通过（%s）：\n  %s\n按 references/01_白板稿.md「故事线」修改 outline.json，再运行 deck storyline <项目>。'
                         % (spath, '\n  '.join(serrs)))
    print('%s（通过）' % spath)
    if '--no-punct' not in sys.argv:
        import punct
        changes, _, mixed = punct.run(proj)
        print('标点整理：去掉 %d 处句号%s' % (len(changes), '；%d 个列表句号不统一，见 00_白板稿/标点整理.md' % len(mixed) if mixed else ''))
    print(render(proj))
    import visual
    errs, warns, plan = visual.run(proj)
    for w in warns:
        print('视觉规划提示：' + w)
    for e in errs:
        print('视觉规划 ✗ ' + e)
    print('%s（%s）' % (plan, '未通过，deck prompts 会拒绝运行' if errs else '通过'))
    print('下一步：把故事线.md 和白板稿发给用户确认（确认点 1），同时用 deck refs 找视觉参考一并询问；用户确认后运行 deck approve。')

if __name__ == '__main__':
    main()

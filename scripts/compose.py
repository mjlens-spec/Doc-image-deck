# -*- coding: utf-8 -*-
"""Stage 3 · Compose the image deck: one full-bleed slide image per page, with the logo(s) and the page number
added as separate PowerPoint objects (never baked into the image), then export a PDF with PowerPoint.

Usage: compose.py <project> [--no-pdf]
Reads  project.json (logos, page_number, upscale), 02_生图/raw/selected.json, 00_白板稿/outline.json
Writes 03_合成/pNN.png (upscaled), 03_合成/<name>_图文版_<suffix>.pptx/.pdf, 03_合成/compose_report.json

Logo entry: {"light": "<logo for light backgrounds>", "dark": "<logo for dark backgrounds>", "corner": "bl|br|tl|tr",
             "height_in": 0.26}. The variant is chosen per slide from the brightness of that corner.
"""
import os, sys, json, argparse
from PIL import Image, ImageStat
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'editable'))
import deckenv as E

W_IN, H_IN = 13.333, 7.5

def corner_stats(img, box):
    g = img.crop(box).convert('L')
    st = ImageStat.Stat(g)
    return st.mean[0], st.stddev[0]

def place(corner, w_in, h_in, margin, foot_y, head_y):
    x = margin if corner.endswith('l') else W_IN - margin - w_in
    y = foot_y if corner.startswith('b') else head_y
    return x, y

def px_box(img, x, y, w, h, pad=0.1):
    W, H = img.size
    return (int(max(0, x - pad) / W_IN * W), int(max(0, y - pad) / H_IN * H),
            int(min(W_IN, x + w + pad) / W_IN * W), int(min(H_IN, y + h + pad) / H_IN * H))

def backing(slide, x, y, w, h):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    pill.name = '页码底'; pill.line.fill.background(); pill.shadow.inherit = False
    pill.fill.solid(); pill.fill.fore_color.rgb = RGBColor(0x10, 0x14, 0x18)
    clr = pill.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
    etree.SubElement(clr, qn('a:alpha')).set('val', '50000')

def main():
    ap = argparse.ArgumentParser(); ap.add_argument('project'); ap.add_argument('--no-pdf', action='store_true')
    a = ap.parse_args()
    proj = os.path.abspath(a.project)
    cfg = E.load_project(proj)
    outline = E.load_outline(proj)
    ids = E.page_ids(outline)
    raw = os.path.join(proj, E.D_GEN, 'raw')
    sel = json.load(open(os.path.join(raw, 'selected.json'), encoding='utf-8'))
    missing = [p for p in ids if p not in sel]
    if missing:
        raise SystemExit('以下页面还没有选定的生图：%s' % ','.join(missing))
    comp = os.path.join(proj, E.D_COMP); os.makedirs(comp, exist_ok=True)
    prs = Presentation(); prs.slide_width, prs.slide_height = E.SW, E.SH
    pn = cfg.get('page_number') or {}
    margin, foot_y, head_y = cfg['margin_in'], cfg.get('footer_y_in', 7.05), cfg.get('header_y_in', 0.2)
    report = []
    for n, (pid, page) in enumerate(zip(ids, outline['pages']), 1):
        src = Image.open(os.path.join(raw, sel[pid])).convert('RGB')
        k = cfg.get('upscale', 2)
        img = src.resize((src.width * k, src.height * k), Image.LANCZOS) if k and k != 1 else src
        png = os.path.join(comp, pid + '.png'); img.save(png)
        jpg = os.path.join(comp, pid + '.jpg'); img.save(jpg, quality=93, subsampling=0, optimize=True)
        media = jpg if os.path.getsize(jpg) < os.path.getsize(png) else png
        s = prs.slides.add_slide(prs.slide_layouts[6])
        pic = s.shapes.add_picture(media, 0, 0, E.SW, E.SH); pic.name = '整页图 %02d' % n
        pic._element._nvXxPr.cNvPr.set('descr', ' '.join(t for _, t in E.page_strings(page))[:1000])
        rec = dict(page=pid, image=sel[pid], media=os.path.basename(media))
        for lg in cfg['logos']:
            h = lg.get('height_in', 0.26)
            path = lg.get('light') or lg.get('dark')
            ar = Image.open(path).size
            w = h * ar[0] / float(ar[1])
            x, y = place(lg.get('corner', 'bl'), w, h, margin, foot_y, head_y)
            mean, _ = corner_stats(img, px_box(img, x, y, w, h))
            dark = mean < 110
            use = (lg.get('dark') if dark else lg.get('light')) or path
            logo = s.shapes.add_picture(use, Inches(x), Inches(y), Inches(w), Inches(h))
            logo.name = lg.get('name', 'Logo')
            rec.setdefault('logos', []).append(('深底用' if dark and lg.get('dark') else '浅底用') + '@' + lg.get('corner', 'bl'))
        first, last = n == 1, n == len(ids)
        if pn.get('corner') and not (first and pn.get('skip_first', True)) and not (last and pn.get('skip_last', True)):
            bw, bh = 0.9, 0.26
            x, y = place(pn['corner'], bw, bh, margin, foot_y, head_y)
            text = pn.get('format', '{:02d}').format(n)
            tw = 0.12 * len(text) + 0.1
            tx = x + bw - tw if pn['corner'].endswith('r') else x
            mean, std = corner_stats(img, px_box(img, tx, y, tw, bh, 0.05))
            dark = mean < 110
            if std > 22:                               # the number would sit on a photo: add a translucent backing
                backing(s, tx - 0.1, y - 0.02, tw + 0.2, bh + 0.04); dark = True
            tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(bw), Inches(bh)); tb.name = '页码'
            tf = tb.text_frame
            for m_ in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
                setattr(tf, m_, 0)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT if pn['corner'].endswith('r') else PP_ALIGN.LEFT
            r = p.add_run(); r.text = text
            r.font.name = pn.get('font', 'Arial'); r.font.size = Pt(pn.get('size_pt', 11))
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if dark else RGBColor(0x5B, 0x6B, 0x79)
            rec['page_number'] = text + (' (底块)' if std > 22 else '')
        if page.get('notes'):
            s.notes_slide.notes_text_frame.text = page['notes']
        report.append(rec)
        os.remove(jpg)                                   # already embedded when used
    prs.core_properties.title = outline.get('title', cfg['name'])
    out = os.path.join(comp, E.out_name(cfg, '图文版', 'pptx'))
    prs.save(out)
    json.dump(report, open(os.path.join(comp, 'compose_report.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(out, '%.1f MB' % (os.path.getsize(out) / 1048576))
    if not a.no_pdf:
        from pptrender import ppt_to_pdf
        pdf = ppt_to_pdf(out, out[:-5] + '.pdf')
        print(pdf)

if __name__ == '__main__':
    main()
